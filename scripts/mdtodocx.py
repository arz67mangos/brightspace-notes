"""Filled course-book .md -> Word: python scripts/mdtodocx.py <in.md> [<out.docx>]
Handles headings, paragraphs, bullets, pipe tables, code blocks and inline **bold**; [[answer]] is shown in blue bold."""
import env, re, sys
from pathlib import Path
from docx import Document
from docx.shared import RGBColor, Pt

TOK = re.compile(r"(\[\[.*?\]\]|\*\*.*?\*\*|\*[^*\s][^*]*?\*)")


def runs(par, text):
    for part in TOK.split(text):
        if not part:
            continue
        if part.startswith("[["):
            r = par.add_run(re.sub(r"\*\*", "", part[2:-2]))
            r.bold = True
            r.font.color.rgb = RGBColor(0x00, 0x40, 0xC0)
        elif part.startswith("**"):
            # answers can sit inside bold: keep the blue styling for [[...]] inside
            for sub in re.split(r"(\[\[.*?\]\])", part[2:-2]):
                if sub.startswith("[["):
                    r = par.add_run(sub[2:-2]); r.bold = True; r.font.color.rgb = RGBColor(0x00, 0x40, 0xC0)
                elif sub:
                    par.add_run(sub).bold = True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            par.add_run(part[1:-1]).italic = True
        else:
            par.add_run(part)


def convert(src, dst):
    doc, lines, i = Document(), Path(src).read_text(encoding="utf-8").splitlines(), 0
    while i < len(lines):
        ln = lines[i]
        if ln.startswith("```"):
            i += 1
            blk = []
            while i < len(lines) and not lines[i].startswith("```"):
                blk.append(lines[i]); i += 1
            p = doc.add_paragraph()
            r = p.add_run("\n".join(blk)); r.font.name = "Consolas"; r.font.size = Pt(9)
        elif ln.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-+:?", c) for c in cells):
                    rows.append(cells)
                i += 1
            i -= 1
            t = doc.add_table(rows=len(rows), cols=max(len(r) for r in rows), style="Table Grid")
            for ri, row in enumerate(rows):
                for ci, c in enumerate(row):
                    cell = t.cell(ri, ci); cell.text = ""
                    runs(cell.paragraphs[0], c)
                    if ri == 0:
                        for r in cell.paragraphs[0].runs:
                            r.bold = True
        elif ln.startswith("#"):
            n = len(ln) - len(ln.lstrip("#"))
            doc.add_heading(ln.lstrip("# ").strip(), level=min(n, 3))
        elif re.match(r"\s*- ", ln):
            lvl = (len(ln) - len(ln.lstrip())) // 2
            runs(doc.add_paragraph(style="List Bullet" if lvl == 0 else "List Bullet 2"), ln.strip()[2:])
        elif ln.strip():
            runs(doc.add_paragraph(), ln.strip())
        i += 1
    doc.save(dst)


if __name__ == "__main__":
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    out = sys.argv[2] if len(sys.argv) > 2 else str(Path(sys.argv[1]).with_suffix(".docx"))
    convert(sys.argv[1], out)
    print("wrote", out)
