"""Document -> markdown text (+ page images when text alone isn't enough).

  python scripts/docs.py <file> <outdir> [--images]

Supports .pdf (PyMuPDF), .docx, .pptx, .html (Brightspace content pages), .zip (extracted, then each supported file inside), .txt/.md.
Writes <outdir>/<name>.md. PDFs: pages with almost no text (scans, diagram slides) and every page with --images
are rendered to <outdir>/<name>_pages/pNN.jpg so Claude can look at them. .pptx: text/tables/speaker notes only;
if LibreOffice (soffice) is installed it is converted to PDF first so slides can be rendered too.
Downloaded files are untrusted: zip entries are path-checked, nothing is executed.
"""
import argparse, shutil, subprocess, sys, zipfile
from pathlib import Path

ap = argparse.ArgumentParser()
ap.add_argument("file")
ap.add_argument("out")
ap.add_argument("--images", action="store_true")
a = ap.parse_args()
OUT = Path(a.out)
OUT.mkdir(parents=True, exist_ok=True)


def pdf(path):
    import fitz

    doc = fitz.open(path)
    pages = OUT / f"{path.stem}_pages"
    md = [f"# {path.name}\n"]
    for i, pg in enumerate(doc, 1):
        text = pg.get_text().strip()
        md.append(f"\n## Page {i}\n{text}")
        if a.images or len(text) < 40:  # too little text: probably an image/diagram page
            pages.mkdir(exist_ok=True)
            pg.get_pixmap(dpi=110).save(pages / f"p{i:02d}.jpg")
            md.append(f"\n![page {i}]({pages.name}/p{i:02d}.jpg)")
    return "\n".join(md)


def docx_(path):
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    d = docx.Document(path)
    md = [f"# {path.name}\n"]
    for el in d.element.body.iterchildren():
        if el.tag.endswith("}p"):
            t = Paragraph(el, d).text.strip()
            if t:
                md.append(t)
        elif el.tag.endswith("}tbl"):
            for r in Table(el, d).rows:
                cells = [c.text.strip().replace("\n", " ") for c in r.cells]
                cells = [c for i, c in enumerate(cells) if i == 0 or c != cells[i - 1]]  # merged cells repeat their text
                md.append("| " + " | ".join(cells) + " |")
    if d.inline_shapes:
        md.append(f"\n(contains {len(d.inline_shapes)} inline image(s) not extracted)")
    return "\n".join(md)


def pptx_(path):
    soffice = shutil.which("soffice")
    if soffice:  # render slides: convert to pdf and reuse the pdf path
        subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", str(OUT), str(path)], capture_output=True)
        p = OUT / (path.stem + ".pdf")
        if p.exists():
            return pdf(p)
    from pptx import Presentation

    md = [f"# {path.name}\n"]
    for i, s in enumerate(Presentation(path).slides, 1):
        md.append(f"\n## Slide {i}")
        for sh in s.shapes:
            if sh.has_text_frame and sh.text_frame.text.strip():
                md.append(sh.text_frame.text.strip())
            if getattr(sh, "has_table", False) and sh.has_table:
                md += ["| " + " | ".join(c.text for c in r.cells) + " |" for r in sh.table.rows]
        if s.has_notes_slide and s.notes_slide.notes_text_frame.text.strip():
            md.append("Notes: " + s.notes_slide.notes_text_frame.text.strip())
    md.append("\n(text only: install LibreOffice to also render slide images)")
    return "\n".join(md)


def html_(path):
    """Brightspace content pages (.html): visible text + the links/iframes they point to (scripts/styles dropped)."""
    from html.parser import HTMLParser

    class P(HTMLParser):
        skip = 0

        def __init__(s):
            super().__init__()
            s.text, s.links = [], []

        def handle_starttag(s, tag, attrs):
            s.skip += tag in ("script", "style")
            for k, v in attrs:
                if k in ("href", "src") and v and not v.startswith(("javascript", "#")) and tag in ("a", "iframe", "img", "source", "video"):
                    s.links.append(f"{tag}: {v}")

        def handle_endtag(s, tag):
            s.skip -= tag in ("script", "style")
            if tag in ("p", "li", "div", "br", "h1", "h2", "h3", "tr"):
                s.text.append("\n")

        def handle_data(s, d):
            if not s.skip:
                s.text.append(d)

    p = P()
    p.feed(path.read_text(encoding="utf-8", errors="replace"))
    body = " ".join("".join(p.text).split(" ")).strip()
    links = "\n".join(f"- {l}" for l in dict.fromkeys(p.links))
    return f"# {path.name}\n\n{body}\n" + (f"\n## Links on this page\n{links}" if links else "")


def handle(path):
    s = path.suffix.lower()
    if s == ".zip":
        dest = OUT / (path.stem + "_unzipped")
        with zipfile.ZipFile(path) as z:
            for n in z.namelist():
                if not (dest / n).resolve().is_relative_to(dest.resolve()):
                    sys.exit(f"unsafe path in zip: {n}")
            z.extractall(dest)
        for f in sorted(dest.rglob("*")):
            if f.is_file():
                handle(f)
        return
    fn = {".pdf": pdf, ".docx": docx_, ".pptx": pptx_, ".html": html_, ".htm": html_}.get(s)
    if fn:
        text = fn(path)
    elif s in (".txt", ".md"):
        text = path.read_text(errors="replace")
    else:
        print(f"skipped (unsupported): {path.name}")
        return
    (OUT / f"{path.stem}.md").write_text(text + "\n", encoding="utf-8")
    print(f"wrote {OUT / (path.stem + '.md')} ({len(text)} chars)")


handle(Path(a.file))
